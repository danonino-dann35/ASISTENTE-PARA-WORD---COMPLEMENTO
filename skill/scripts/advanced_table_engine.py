# -*- coding: utf-8 -*-
"""
Motor de Formato y Manipulación Avanzada de Tablas para Word (.docx)
Permite construir y transformar tablas con formato estricto APA 7ma edición,
combinación de celdas (merge), sombreados personalizados y alineaciones.
"""

import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def apply_apa7_borders(table):
    """Aplica el estándar visual estricto de bordes APA 7ma edición."""
    tblPr = table._tbl.tblPr
    # Remover bordes previos
    for child in list(tblPr):
        if child.tag.endswith('tblBorders'):
            tblPr.remove(child)

    # Definir bordes: superior simple, inferior simple, sin bordes verticales
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:insideH w:val="none"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

    # Agregar línea horizontal inferior a la fila de encabezado
    if len(table.rows) > 0:
        header_tr = table.rows[0]._tr
        trPr = header_tr.get_or_add_trPr()
        # Aplicar borde inferior en las celdas del encabezado
        for cell in table.rows[0].cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>\n'
                f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
                f'</w:tcBorders>'
            )
            tcPr.append(tcBorders)

def set_cell_shading(cell, color_hex):
    """Aplica color de fondo (sombreado) a una celda específica."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def merge_cells(table, row1, col1, row2, col2):
    """Combina un rango de celdas dentro de una tabla."""
    cell_start = table.cell(row1, col1)
    cell_end = table.cell(row2, col2)
    return cell_start.merge(cell_end)

def create_formatted_table(doc, headers, data_rows, title=None, note=None, apa7=True):
    """Crea e inserta una tabla completamente formateada en el documento."""
    if title:
        p_num = doc.add_paragraph()
        p_num.paragraph_format.space_before = Pt(12)
        p_num.paragraph_format.space_after = Pt(2)
        r_num = p_num.add_run("Tabla")
        r_num.font.name = "Times New Roman"
        r_num.font.bold = True

        p_tit = doc.add_paragraph()
        p_tit.paragraph_format.space_before = Pt(0)
        p_tit.paragraph_format.space_after = Pt(6)
        r_tit = p_tit.add_run(title)
        r_tit.font.name = "Times New Roman"
        r_tit.font.italic = True

    total_rows = len(data_rows) + 1
    total_cols = len(headers)
    table = doc.add_table(rows=total_rows, cols=total_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Encabezados
    for col_idx, h_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = str(h_text)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run.font.bold = True

    # Filas de datos
    for r_idx, row_values in enumerate(data_rows):
        for c_idx, val in enumerate(row_values):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = str(val)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(10)

    if apa7:
        apply_apa7_borders(table)

    if note:
        p_note = doc.add_paragraph()
        p_note.paragraph_format.space_before = Pt(4)
        p_note.paragraph_format.space_after = Pt(12)
        r_nlbl = p_note.add_run("Nota. ")
        r_nlbl.font.name = "Times New Roman"
        r_nlbl.font.size = Pt(10)
        r_nlbl.font.italic = True
        r_ntxt = p_note.add_run(note)
        r_ntxt.font.name = "Times New Roman"
        r_ntxt.font.size = Pt(10)

    return table

if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--test":
        doc_test = docx.Document()
        h = ["Dimensión", "Pretest", "Postest", "Diferencia"]
        d = [
            ["Fuerza Vertical", "4.41", "6.44", "+2.03"],
            ["Fuerza Horizontal", "4.31", "6.03", "+1.72"]
        ]
        create_formatted_table(doc_test, h, d, title="Rendimiento Comparativo", note="Datos de prueba.", apa7=True)
        print("Motor de tablas avanzadas validado con éxito.")
